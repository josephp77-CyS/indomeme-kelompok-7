"""Create editable DOCX and standalone HTML from the reviewed plain-language write-up."""
import base64
import hashlib
import html
import json
from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor
from PIL import Image
from math_nodes import word_math, browser_math
from equations import eqs,inline_symbols,inline_tex

WORK=Path(__file__).resolve().parents[2]
OUT=WORK/'docs'
STEM='laporan'
HERE=Path(__file__).parent
source=(HERE/'source.md').read_text(encoding='utf-8')
formulas={e['id']:e for e in json.loads((HERE/'formulas.json').read_text(encoding='utf-8'))}
INLINE_RE=re.compile(r'(?<![\w^])('+'|'.join(re.escape(k) for k in sorted(inline_symbols,key=len,reverse=True))+r')(?![\w^])')


def font(run,size=None,bold=None,italic=None):
    run.font.name='Arial'
    prop=run._element.get_or_add_rPr();fonts=prop.get_or_add_rFonts()
    fonts.set(qn('w:ascii'),'Arial');fonts.set(qn('w:hAnsi'),'Arial')
    run.font.color.rgb=RGBColor(0,0,0)
    if size is not None:run.font.size=Pt(size)
    if bold is not None:run.bold=bold
    if italic is not None:run.italic=italic


def add_text(paragraph,text,size=None):
    pieces=re.split(r'(\*\*.*?\*\*|\*[^*]+\*)',text)
    for piece in pieces:
        if not piece:continue
        bold=piece.startswith('**') and piece.endswith('**')
        italic=not bold and piece.startswith('*') and piece.endswith('*')
        plain=piece[2:-2] if bold else piece[1:-1] if italic else piece
        for segment in re.split(r'(https?://\S+)',plain):
            if segment.startswith(('https://','http://')):
                url=segment.rstrip('.')
                relationship=paragraph.part.relate_to(url,RT.HYPERLINK,is_external=True)
                link=OxmlElement('w:hyperlink');link.set(qn('r:id'),relationship)
                run=OxmlElement('w:r');prop=OxmlElement('w:rPr')
                color=OxmlElement('w:color');color.set(qn('w:val'),'24566F');prop.append(color)
                run.append(prop);value=OxmlElement('w:t');value.text=url;run.append(value);link.append(run);paragraph._p.append(link)
                if segment.endswith('.'):font(paragraph.add_run('.'),size,bold,italic)
            else:
                for value in INLINE_RE.split(segment):
                    if value in inline_symbols:paragraph._p.append(word_math(row_math(value)))
                    else:font(paragraph.add_run(value),size,bold,italic)


def inline_html(text):
    value=html.escape(text)
    value=re.sub(r'\*\*(.+?)\*\*',r'<strong>\1</strong>',value)
    value=re.sub(r'(?<!\*)\*([^*]+)\*(?!\*)',r'<em>\1</em>',value)
    value=re.sub(r'(https?://[^\s<]+)',lambda m:'<a href="'+m.group(1).rstrip('.')+'">'+m.group(1).rstrip('.')+'</a>'+('.' if m.group(1).endswith('.') else ''),value)
    return INLINE_RE.sub(lambda m:browser_math(row_math(m.group(0)),m.group(0),False),value)


def row_math(key):
    from math_nodes import row
    return row(inline_symbols[key])


def table(doc,header,rows):
    t=doc.add_table(rows=1,cols=len(header));t.autofit=False
    if len(header)==5:widths=[.65,1.15,2.1,1.25,1.25]
    elif len(header)==3 and header[0]=='Bagian data':widths=[1.4,.9,4.1]
    elif len(header)==3 and header[0]=='Varian':widths=[2.8,1.8,1.8]
    elif len(header)==3:widths=[2.0,2.2,2.2]
    else:widths=[1.4,5.0] if header[0]=='Istilah' else [2.6,3.8]
    t.alignment=1
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    for i,head in enumerate(header):t.rows[0].cells[i].text=head
    for row in rows:
        cells=t.add_row().cells
        for i,item in enumerate(row):cells[i].text=item
    for ri,row in enumerate(t.rows):
        for ci,cell in enumerate(row.cells):
            cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            prop=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:color'),'D9D9D9');e.set(qn('w:sz'),'4');borders.append(e)
            prop.append(borders)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E5ECF0' if ri==0 else 'FFFFFF');prop.append(shade)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');margins.append(e)
            prop.append(margins)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.05
                for run in p.runs:font(run,10.1,bold=ri==0)
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))


def parse(markdown):
    items=[]
    blocks=markdown.strip().split('\n\n')
    for block in blocks:
        if block.startswith('# '):items.append(('title',block[2:]))
        elif block.startswith('### '):items.append(('subheading',block[4:]))
        elif block.startswith('## '):items.append(('heading',block[3:]))
        elif block.startswith('Rumus: '):items.append(('equation',block[7:].strip()))
        elif block.startswith('!['):
            match=re.fullmatch(r'!\[(.*?)\]\((.*?)\)',block)
            assert match,block
            items.append(('figure',match[1],match[2]))
        elif block.startswith('| '):
            rows=[[x.strip() for x in line.strip().strip('|').split('|')] for line in block.splitlines()]
            assert all(x=='---' for x in rows[1])
            items.append(('table',rows[0],rows[2:]))
        elif block.startswith('- '):
            for line in block.splitlines():items.append(('bullet',line[2:]))
        elif block.startswith('*Gambar '):items.append(('caption',block.strip('*')))
        else:items.append(('paragraph',block.replace('\n',' ')))
    return items


def build():
    items=parse(source)
    assert len([i for i in items if i[0]=='figure'])==4
    equation_ids=[i[1] for i in items if i[0]=='equation']
    assert len(equation_ids)>=14 and len(equation_ids)==len(set(equation_ids))
    assert all(key in eqs and key in formulas for key in equation_ids)
    doc=Document();section=doc.sections[0]
    section.page_width=Inches(8.5);section.page_height=Inches(11)
    section.top_margin=Inches(.72);section.bottom_margin=Inches(.7)
    section.left_margin=Inches(.9);section.right_margin=Inches(.9)
    for name in ['Normal','Title','Heading 1','Heading 2','Caption','List Bullet']:
        doc.styles[name].font.name='Arial';doc.styles[name].font.color.rgb=RGBColor(0,0,0)
    normal=doc.styles['Normal'];normal.font.size=Pt(11)
    normal.paragraph_format.line_spacing=1.1;normal.paragraph_format.space_after=Pt(5)
    normal.paragraph_format.keep_together=True
    doc.styles['Title'].font.size=Pt(24);doc.styles['Title'].font.bold=True
    doc.styles['Title'].paragraph_format.space_after=Pt(10)
    heading=doc.styles['Heading 1'];heading.font.size=Pt(15)
    heading.paragraph_format.space_before=Pt(13);heading.paragraph_format.space_after=Pt(6)
    heading.paragraph_format.keep_with_next=True
    doc.styles['Heading 2'].font.size=Pt(12)
    doc.styles['Heading 2'].paragraph_format.keep_with_next=True
    doc.styles['Heading 2'].paragraph_format.space_before=Pt(9)
    caption=doc.styles['Caption'];caption.font.size=Pt(9.3);caption.font.bold=False
    caption.paragraph_format.space_before=Pt(2);caption.paragraph_format.space_after=Pt(8)
    doc.core_properties.title='Reproduksi Eksperimen Inti INDOMEME'
    doc.core_properties.author='Kelompok 7'
    doc.core_properties.subject='Laporan lengkap dengan rumusan teknis dan penjelasan bertahap'
    toc=[(f'bagian-{i}',item[1]) for i,item in enumerate(items) if item[0]=='heading']
    parts=[];figure_records=[];paragraph_count=0;equation_count=0
    for item in items:
        kind=item[0]
        if kind=='title':
            doc.add_paragraph(item[1],style='Title');parts.append('<h1>'+inline_html(item[1])+'</h1>')
        elif kind in ['heading','subheading']:
            level=1 if kind=='heading' else 2
            doc.add_paragraph(item[1],style='Heading '+str(level))
            anchor=next((key for key,value in toc if value==item[1]),'')
            tag='h2' if level==1 else 'h3'
            parts.append('<'+tag+' id="'+anchor+'">'+inline_html(item[1])+'</'+tag+'>')
        elif kind=='equation':
            equation_count+=1;key=item[1]
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_together=True;p.paragraph_format.keep_with_next=True
            p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(7)
            p._p.append(word_math(eqs[key]));font(p.add_run('  ('+str(equation_count)+')'),9)
            parts.append('<div class="equation">'+browser_math(eqs[key],formulas[key]['alt'])+'<span>('+str(equation_count)+')</span></div>')
        elif kind in ['paragraph','bullet','caption']:
            style='List Bullet' if kind=='bullet' else 'Caption' if kind=='caption' else None
            p=doc.add_paragraph(style=style)
            add_text(p,item[1],9.3 if kind=='caption' else None)
            if kind=='caption':p.paragraph_format.keep_together=True
            if kind=='paragraph':
                paragraph_count+=1
                if paragraph_count<=2:
                    for r in p.runs:font(r,10)
                elif paragraph_count==3:
                    for r in p.runs:font(r,12,bold=True)
            cls='caption' if kind=='caption' else 'meta' if paragraph_count<=2 and kind=='paragraph' else 'lead' if paragraph_count==3 and kind=='paragraph' else 'body'
            parts.append('<p class="'+cls+'">'+inline_html(item[1])+'</p>')
        elif kind=='figure':
            image_path=(OUT/item[2]).resolve()
            assert image_path.is_relative_to(OUT) and image_path.is_file(),image_path
            with Image.open(image_path) as image:width,height=image.size
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(2)
            p.paragraph_format.keep_with_next=True
            run=p.add_run();run.add_picture(str(image_path),width=Inches(6.4))
            drawing=run._r.xpath('.//wp:docPr')[0];drawing.set('descr',item[1])
            encoded=base64.b64encode(image_path.read_bytes()).decode()
            parts.append('<figure><img src="data:image/png;base64,'+encoded+'" alt="'+html.escape(item[1])+'"></figure>')
            figure_records.append({'file':item[2],'pixels':[width,height],'sha256':hashlib.sha256(image_path.read_bytes()).hexdigest()})
        elif kind=='table':
            table(doc,item[1],item[2]);parts.append('<table><thead><tr>'+''.join('<th>'+inline_html(x)+'</th>' for x in item[1])+'</tr></thead><tbody>'+''.join('<tr>'+''.join('<td>'+inline_html(x)+'</td>' for x in row)+'</tr>' for row in item[2])+'</tbody></table>')
    footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    font(footer.add_run('Kelompok 7  '),8.5)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    for element in [doc._element,doc.styles.element]:
        for border in element.xpath('.//w:pBdr'):border.getparent().remove(border)
    doc.save(OUT/(STEM+'.docx'))
    css='''
    :root{color-scheme:light}*{box-sizing:border-box}body{margin:0;background:#fff;color:#1c2b35;font:18px/1.75 system-ui,-apple-system,BlinkMacSystemFont,"Segoe UI",sans-serif}
    article{max-width:940px;margin:48px auto 72px;padding:0 30px}h1{font-size:42px;line-height:1.16;letter-spacing:-.025em;margin:0 0 22px;color:#000}h2{font-size:27px;line-height:1.3;margin:46px 0 14px;color:#000}
    p{margin:0 0 18px}.meta{font-size:14px;color:#53636d;margin-bottom:8px}.lead{font-size:22px;line-height:1.5;font-weight:650;margin-top:22px}strong{font-weight:680}figure{margin:28px 0 6px}img{display:block;width:100%;height:auto}.caption{font-size:14px;line-height:1.55;color:#465a67;margin:4px 0 28px}
    a{color:#1b607e;text-underline-offset:3px;overflow-wrap:anywhere}table{width:100%;border-collapse:collapse;font-size:16px;margin:20px 0 28px}th,td{border:1px solid #d9d9d9;padding:10px 13px;vertical-align:top;text-align:left}th{background:#e5ecf0;color:#142b3a}td:first-child{width:145px}
    @media(max-width:600px){article{padding:0 20px;margin-top:28px}h1{font-size:32px}h2{font-size:23px}body{font-size:17px}table{font-size:14px}td:first-child{width:105px}}@media print{article{max-width:none;margin:0;padding:0}h2{break-after:avoid}figure{break-inside:avoid}a{color:inherit}}
    '''
    css+='h3{font-size:21px;line-height:1.35;margin:26px 0 12px;color:#000}.equation{display:flex;align-items:center;justify-content:center;gap:18px;margin:22px 0;overflow-x:auto;padding:8px 0}.equation math{font-size:1.1em;flex:0 0 auto}.equation span{font-size:14px;color:#53636d}details{margin:22px 0 30px}summary{cursor:pointer;font-weight:650}nav a{display:block;font-size:15px;line-height:1.8}thead{display:table-header-group}tr{break-inside:avoid}'
    nav='<details><summary>Daftar isi</summary><nav>'+''.join('<a href="#'+key+'">'+html.escape(value)+'</a>' for key,value in toc)+'</nav></details>'
    parts.insert(3,nav)
    page='<!doctype html><html lang="id"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Reproduksi Eksperimen Inti INDOMEME</title><style>'+css+'</style></head><body><article>'+''.join(parts)+'</article></body></html>'
    (OUT/(STEM+'.html')).write_text(page,encoding='utf-8')
    published=INLINE_RE.sub(lambda m:'$'+inline_tex[m.group(0)]+'$',source)
    for key in equation_ids:published=published.replace('Rumus: '+key+'\n','$$\n'+formulas[key]['latex']+'\n$$\n')
    assert 'Rumus: ' not in published
    (OUT/(STEM+'.md')).write_text(published,encoding='utf-8')
    manifest={'word_count_source':len(source.split()),'figures':figure_records,'equation_ids':equation_ids,
      'output_files':{name:hashlib.sha256((OUT/(STEM+name)).read_bytes()).hexdigest() for name in ['.md','.docx','.html']}}
    (Path(__file__).parent/'build_manifest.json').write_text(json.dumps(manifest,indent=2))
    assert len(doc._element.xpath('.//m:oMath'))>=equation_count
    assert page.count('display="block"')==equation_count
    print(json.dumps({'word_count':manifest['word_count_source'],'figures':len(figure_records),'native_equations':equation_count,'outputs':list(manifest['output_files'])}))


if __name__=='__main__':build()
